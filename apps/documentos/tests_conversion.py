"""Tests del tamaño de papel, de los márgenes y del saneo de la copia al
convertir documentos office a PDF y de la regeneración automática cuando el
PDF del disco desaparece.

El PDF hereda el ``pageSetup`` del archivo de origen. Si el origen es A4 y el
contenido es más alto que la A4 imprimible, las tablas se parten entre páginas;
si el papel es demasiado alto (Folio), el contenido de la página siguiente se
descoloca con el de la anterior. Por eso se parchea en una copia temporal el
tamaño de papel (por defecto **A4**) y los **márgenes superior/inferior**
(0.4", que es lo que permite que el formulario siga cabiendo en una sola
página): el archivo subido nunca se modifica.

Esa misma copia se **sanea** forzando el salto de página del título repetido
que dejaba el encabezado de la sección de instrucciones huérfano. Cada
sección tiene su propio logo (el de la 1ª en ``A1:A5``, el de la 2ª en
``A60:A64``) y **ninguno se retira**: el de la última página desapareció
cuando se borró "por duplicado".

Además ``Documento.archivo_pdf`` guarda una ruta en BD que puede quedar
"fantasma" si borran el PDF del disco: aquí se comprueba que las vistas
(``ver_archivo``, ``detalle``, ``estado_conversion``, ``descargar_archivo``, preview de
planilla) no den 500 y regeneren el PDF en línea.
"""

import base64
import os
import re
import shutil
import subprocess
import tempfile
import unittest
import zipfile
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, override_settings
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ImagenXlsx
from openpyxl.worksheet.pagebreak import Break

from tests_base import BaseDocumentosTest
from documentos.models import Documento, PlanillaCambios
from documentos.services.conversion_pdf import (
    convertir_office_a_pdf,
    parchear_tamano_papel,
    resolver_margenes,
    resolver_papel,
)
from documentos.services.documentos_pdf import (
    archivo_existe,
    asegurar_pdf,
    convertir_documento_sync,
    limpiar_pdf_inexistente,
    pdf_disponible,
)
from documentos.tasks import (
    convertir_a_pdf_sync,
    convertir_documento_a_pdf,
    convertir_planilla_a_pdf,
)

NS_XLSX = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
NS_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'

LIBREOFFICE = shutil.which('libreoffice')


def paper_sizes(ruta):
    """`paperSize` declarados en las hojas de un XLSX."""
    with zipfile.ZipFile(ruta) as z:
        tallas = []
        for nombre in sorted(z.namelist()):
            if re.match(r'xl/worksheets/sheet\d+\.xml$', nombre):
                tallas += re.findall(
                    r'paperSize="(\d+)"',
                    z.read(nombre).decode('utf-8', 'replace'),
                )
        return tallas


def pg_sz_docx(ruta):
    """`(w, h, orient)` de la primera sección de un DOCX."""
    with zipfile.ZipFile(ruta) as z:
        xml = z.read('word/document.xml').decode('utf-8', 'replace')
    w = re.search(r'w:pgSz[^>]*w:w="(\d+)"', xml)
    h = re.search(r'w:pgSz[^>]*w:h="(\d+)"', xml)
    orient = re.search(r'w:pgSz[^>]*w:orient="(\w+)"', xml)
    return (
        int(w.group(1)) if w else None,
        int(h.group(1)) if h else None,
        orient.group(1) if orient else None,
    )


def margenes_xlsx(ruta, hoja='xl/worksheets/sheet1.xml'):
    """`(top, bottom, left)` en pulgadas de un XLSX, o None si no hay."""
    with zipfile.ZipFile(ruta) as z:
        xml = z.read(hoja).decode('utf-8', 'replace')
    coincide = re.search(r'<pageMargins[^>]*/>', xml)
    if not coincide:
        return None
    etiqueta = coincide.group(0)
    valores = []
    for atributo in ('top', 'bottom', 'left'):
        m = re.search(r'%s="([\d\.]+)"' % atributo, etiqueta)
        valores.append(float(m.group(1)) if m else None)
    return tuple(valores)


def pg_mar_docx(ruta):
    """`(top, bottom, left)` en twips de un DOCX, o None si no hay."""
    with zipfile.ZipFile(ruta) as z:
        xml = z.read('word/document.xml').decode('utf-8', 'replace')
    coincide = re.search(r'<w:pgMar[^>]*/>', xml)
    if not coincide:
        return None
    etiqueta = coincide.group(0)
    valores = []
    for atributo in ('top', 'bottom', 'left'):
        m = re.search(r'w:%s="(-?\d+)"' % atributo, etiqueta)
        valores.append(int(m.group(1)) if m else None)
    return tuple(valores)


def primera_mediabox(ruta):
    """`(ancho, alto)` en puntos de la primera página de un PDF."""
    with open(ruta, 'rb') as f:
        datos = f.read()
    match = re.search(
        rb'MediaBox\s*\[\s*[\d\.\-]+\s+[\d\.\-]+\s+([\d\.]+)\s+([\d\.]+)\s*\]',
        datos,
    )
    if not match:
        raise AssertionError('Sin MediaBox en %s' % ruta)
    return float(match.group(1)), float(match.group(2))


def crear_xlsx(ruta, hojas):
    """XLSX mínimo. `hojas` = lista de `<pageSetup .../>` (o None por hoja)."""
    with zipfile.ZipFile(ruta, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml', '<?xml version="1.0"?><Types/>')
        for i, page_setup in enumerate(hojas, start=1):
            z.writestr(
                'xl/worksheets/sheet%d.xml' % i,
                (
                    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                    '<worksheet %s>'
                    '<sheetData><row r="1"><c r="A1" t="inlineStr">'
                    '<is><t>Hola</t></is></c></row></sheetData>'
                    '%s'
                    '</worksheet>'
                ) % (NS_XLSX, page_setup or ''),
            )
    return ruta


def crear_docx(ruta, w, h, orient=None):
    """DOCX mínimo (suficiente: el parcheo solo lee `word/document.xml`)."""
    attrs = 'w:w="%d" w:h="%d"' % (w, h)
    if orient:
        attrs += ' w:orient="%s"' % orient
    xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document %s><w:body><w:p/><w:sectPr>'
        '<w:pgSz %s/>'
        '<w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440"/>'
        '</w:sectPr></w:body></w:document>'
    ) % (NS_W, attrs)
    with zipfile.ZipFile(ruta, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('word/document.xml', xml)
    return ruta


def xlsx_real(ruta, paper_size=9):
    """XLSX válido (openpyxl) con `paperSize` declarado."""
    wb = Workbook()
    ws = wb.active
    ws['A1'] = 'Formulario de prueba'
    ws['A2'] = 'Fila 2'
    ws.page_setup.paperSize = paper_size
    ws.page_setup.orientation = 'portrait'
    wb.save(ruta)
    return ruta


def xlsx_de_filas(ruta, filas=50, alto_fila=15, margen=0.75, paper_size=9):
    """XLSX con `filas` filas de `alto_fila` puntos y márgenes `margen` (pulg).

    Útil para comprobar empíricamente cuánto contenido cabe en una página.
    """
    wb = Workbook()
    ws = wb.active
    for i in range(1, filas + 1):
        ws.cell(row=i, column=1, value='Fila %d' % i)
        ws.row_dimensions[i].height = alto_fila
    ws.page_setup.paperSize = paper_size
    ws.page_setup.orientation = 'portrait'
    ws.page_margins.top = margen
    ws.page_margins.bottom = margen
    ws.page_margins.header = 0.3
    ws.page_margins.footer = 0.3
    wb.save(ruta)
    return ruta


def contar_paginas(ruta):
    """Número de páginas de un PDF sin depender de una librería."""
    with open(ruta, 'rb') as f:
        datos = f.read()
    return len(re.findall(rb'/Type\s*/Page(?!s)', datos))


#: PNG de 16×16 px para los tests que necesitan una imagen real en el XLSX.
PNG_16 = base64.b64decode(
    'iVBORw0KGgoAAAANSUhEUgAAABAAAAAQCAIAAACQkWg2AAAAGUlEQVR4nGNkWBbIQApgIkn1'
    'qIZRDUNKAwAmbgEXSZwHbQAAAABJRU5ErkJggg=='
)


def _inyectar_logo_encabezado(ruta):
    """Añade a ``ruta`` el logo de cabecera (``twoCellAnchor`` en A1).

    Reutiliza el ``rId`` del logo que openpyxl ya guardó, igual que las
    plantillas SGC (ambos logos comparten la misma imagen). El dibujo debe
    existir previamente en el XLSX.
    """
    with zipfile.ZipFile(ruta) as z:
        partes = {nombre: z.read(nombre) for nombre in z.namelist()}
    dibujo = partes['xl/drawings/drawing1.xml'].decode('utf-8')
    rid = re.search(r'r:embed="([^"]+)"', dibujo).group(1)
    ancla = (
        '<twoCellAnchor>'
        '<from><col>0</col><colOff>0</colOff><row>0</row><rowOff>0</rowOff></from>'
        '<to><col>2</col><colOff>0</colOff><row>2</row><rowOff>0</rowOff></to>'
        '<pic><nvPicPr><cNvPr id="99" name="Logo encabezado"/><cNvPicPr/></nvPicPr>'
        '<blipFill><a:blip xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'cstate="print" r:embed="{0}"/>'
        '<a:stretch xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<a:fillRect/></a:stretch></blipFill>'
        '<spPr><a:prstGeom xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'prst="rect"/></spPr></pic><clientData/></twoCellAnchor>'
    ).format(rid)
    inicio = dibujo.index('>') + 1
    partes['xl/drawings/drawing1.xml'] = (
        dibujo[:inicio] + ancla + dibujo[inicio:]
    ).encode('utf-8')
    with zipfile.ZipFile(ruta, 'w', zipfile.ZIP_DEFLATED) as z:
        for nombre, contenido in partes.items():
            z.writestr(nombre, contenido)
    return ruta


def xlsx_secciones_con_logo(ruta, filas=40, repite_en=30, encabezado=True):
    """XLSX con dos secciones y un logo por sección.

    Reproduce el patrón de las plantillas SGC-FOR-001/002/005: el título se
    repite en ``repite_en`` (la 2ª sección) y en esa misma fila se ancla su
    logo (``oneCellAnchor``), mientras que con ``encabezado`` se inyecta además
    el logo de la 1ª sección (``twoCellAnchor`` en A1). **Ninguno de los dos
    logos se debe retirar**: el de la 2ª sección es el de la última página.
    Con ``encabezado=False`` solo queda el logo suelto, que además no se toca.
    """
    logo = os.path.join(tempfile.gettempdir(), 'polidoc_logo_test.png')
    with open(logo, 'wb') as f:
        f.write(PNG_16)

    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value='TITULO DE PRUEBA')
    ws.row_dimensions[1].height = 15
    for i in range(2, filas + 1):
        ws.cell(row=i, column=1,
                value='TITULO DE PRUEBA' if i == repite_en else 'Fila %d' % i)
        ws.row_dimensions[i].height = 15
    ws.add_image(ImagenXlsx(logo), 'A%d' % repite_en)
    ws.page_setup.paperSize = 9
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4
    wb.save(ruta)

    if encabezado:
        return _inyectar_logo_encabezado(ruta)
    return ruta


def xlsx_con_titulo_repetido(ruta, filas=40, repite_en=30,
                             titulo='TITULO DE PRUEBA', salto_previo=None):
    """XLSX que repite su primera fila (el título) más abajo.

    ``repite_en`` admite un nº de fila o una secuencia; con ``None`` no hay
    repetición. ``salto_previo`` añade ya un ``<rowBreaks>`` a mano.
    """
    if repite_en is None:
        repetidas = ()
    elif isinstance(repite_en, int):
        repetidas = (repite_en,)
    else:
        repetidas = tuple(repite_en)

    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value=titulo)
    ws.row_dimensions[1].height = 15
    for i in range(2, filas + 1):
        ws.cell(row=i, column=1,
                value=titulo if i in repetidas else 'Fila %d' % i)
        ws.row_dimensions[i].height = 15
    ws.page_setup.paperSize = 9
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4
    if salto_previo is not None:
        ws.row_breaks.append(Break(id=salto_previo))
    wb.save(ruta)
    return ruta


def anclas_dibujo(ruta, dibujo='xl/drawings/drawing1.xml'):
    """Tipos de ancla declarados en un dibujo del XLSX."""
    with zipfile.ZipFile(ruta) as z:
        xml = z.read(dibujo).decode('utf-8', 'replace')
    return re.findall(r'<(?:\w+:)?(twoCellAnchor|oneCellAnchor|absoluteAnchor)[ />]', xml)


def filas_con_salto(ruta, hoja='xl/worksheets/sheet1.xml'):
    """``id`` de los ``<brk>`` de una hoja (lista vacía si no hay)."""
    with zipfile.ZipFile(ruta) as z:
        xml = z.read(hoja).decode('utf-8', 'replace')
    bloque = re.search(r'<rowBreaks[^>]*>(.*?)</rowBreaks>', xml, re.S)
    if not bloque:
        return []
    return [int(i) for i in re.findall(r'<brk[^>]*id="(\d+)"', bloque.group(1))]


def paginas_texto(ruta):
    """Texto de cada página de un PDF."""
    from pypdf import PdfReader
    return [(pagina.extract_text() or '') for pagina in PdfReader(ruta).pages]


def dibujos_pdf(ruta):
    """``(página, matriz, recurso)`` de cada imagen dibujada en el PDF."""
    from pypdf import PdfReader
    salidas = []
    for numero, pagina in enumerate(PdfReader(ruta).pages, start=1):
        datos = pagina.get_contents().get_data().decode('latin-1')
        for m in re.finditer(
            r'([-\d\.]+ [-\d\.]+ [-\d\.]+ [-\d\.]+ [-\d\.]+ [-\d\.]+) cm', datos,
        ):
            operador = re.match(r'\s*/(\w+) Do', datos[m.end():m.end() + 200])
            if operador:
                salidas.append((numero, m.group(1).strip(), operador.group(1)))
    return salidas


class BaseTemporalTest(unittest.TestCase):
    """Directorio temporal por test."""

    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix='polidoc-conv-')
        self.addCleanup(shutil.rmtree, self.dir, ignore_errors=True)

    def ruta(self, nombre):
        return os.path.join(self.dir, nombre)


# =============================================================================
# RESOLUCIÓN DEL PAPEL
# =============================================================================

class ResolverPapelTest(BaseTemporalTest):

    def test_folio(self):
        self.assertEqual(resolver_papel('Folio'), (14, 12240, 18720))

    def test_oficio_es_folio(self):
        self.assertEqual(resolver_papel('oficio'), (14, 12240, 18720))

    def test_legal(self):
        self.assertEqual(resolver_papel('Legal'), (5, 12240, 20160))

    def test_carta(self):
        self.assertEqual(resolver_papel('carta'), (1, 12240, 15840))

    def test_a4(self):
        codigo, ancho, alto = resolver_papel('A4')
        self.assertEqual(codigo, 9)
        self.assertGreater(ancho, 0)
        self.assertGreater(alto, ancho)

    def test_codigo_numerico(self):
        self.assertEqual(resolver_papel('14'), (14, 12240, 18720))

    def test_vacio_no_parchea(self):
        self.assertEqual(resolver_papel(''), (None, None, None))
        self.assertEqual(resolver_papel('original'), (None, None, None))

    def test_desconocido_no_parchea(self):
        self.assertEqual(resolver_papel('tamaño-inexistente'), (None, None, None))

    @override_settings(PAPEL_PDF_CONVERSION='Legal')
    def test_usa_el_setting_por_defecto(self):
        self.assertEqual(resolver_papel(None), (5, 12240, 20160))

    def test_por_defecto_es_a4(self):
        """A4: el oficio hacía páginas tan altas que se descolocaban."""
        codigo, ancho, alto = resolver_papel(None)
        self.assertEqual(codigo, 9)
        self.assertEqual((ancho, alto), (int(8.27 * 1440), int(11.69 * 1440)))


class ResolverMargenesTest(unittest.TestCase):

    @override_settings(MARGEN_SUPERIOR_PDF_CONVERSION=0.4,
                       MARGEN_INFERIOR_PDF_CONVERSION=0.4)
    def test_por_defecto_son_0_4(self):
        self.assertEqual(resolver_margenes(), {'top': 0.4, 'bottom': 0.4})

    @override_settings(MARGEN_SUPERIOR_PDF_CONVERSION=None,
                       MARGEN_INFERIOR_PDF_CONVERSION=None)
    def test_vacios_no_tocan_nada(self):
        self.assertIsNone(resolver_margenes())

    @override_settings(MARGEN_SUPERIOR_PDF_CONVERSION='',
                       MARGEN_INFERIOR_PDF_CONVERSION='')
    def test_cadena_vacia_no_toca_nada(self):
        self.assertIsNone(resolver_margenes())

    def test_solo_un_lado(self):
        self.assertEqual(resolver_margenes({'top': 0.3, 'bottom': None}),
                         {'top': 0.3})

    def test_cadenas_y_decimales(self):
        self.assertEqual(resolver_margenes({'top': '0.25', 'bottom': 5}),
                         {'top': 0.25, 'bottom': 5.0})

    def test_valores_invalidos_se_ignoran(self):
        self.assertIsNone(resolver_margenes({'top': 'nada', 'bottom': -1}))


# =============================================================================
# PARCHEO DE LA COPIA
# =============================================================================

class ParcheoXlsxTest(BaseTemporalTest):

    def test_cambia_a_folio_sin_tocar_el_original(self):
        origen = crear_xlsx(self.ruta('origen.xlsx'), ['<pageSetup orientation="portrait" paperSize="9"/>'])

        copia, temp_dir = parchear_tamano_papel(origen, 'Folio')

        self.assertNotEqual(copia, origen)
        self.assertTrue(os.path.exists(copia))
        self.assertEqual(paper_sizes(copia), ['14'])
        self.assertEqual(paper_sizes(origen), ['9'])
        if temp_dir:
            shutil.rmtree(temp_dir, ignore_errors=True)

    def test_todas_las_hojas(self):
        origen = crear_xlsx(self.ruta('multi.xlsx'), [
            '<pageSetup paperSize="9"/>',
            '<pageSetup paperSize="9"/>',
        ])

        copia, temp_dir = parchear_tamano_papel(origen, 'Folio')

        self.assertEqual(paper_sizes(copia), ['14', '14'])
        if temp_dir:
            shutil.rmtree(temp_dir, ignore_errors=True)

    def test_inserta_page_setup_si_falta(self):
        origen = crear_xlsx(self.ruta('sin.xlsx'), [None])

        copia, temp_dir = parchear_tamano_papel(origen, 'Folio')

        self.assertEqual(paper_sizes(copia), ['14'])
        if temp_dir:
            shutil.rmtree(temp_dir, ignore_errors=True)

    def test_conserva_orientacion_apaisada(self):
        origen = crear_xlsx(
            self.ruta('pais.xlsx'),
            ['<pageSetup orientation="landscape" paperSize="9"/>'],
        )

        copia, temp_dir = parchear_tamano_papel(origen, 'Folio')

        with zipfile.ZipFile(copia) as z:
            xml = z.read('xl/worksheets/sheet1.xml').decode('utf-8', 'replace')
        self.assertIn('paperSize="14"', xml)
        self.assertIn('orientation="landscape"', xml)
        if temp_dir:
            shutil.rmtree(temp_dir, ignore_errors=True)

    def test_xlsx_corrupto_no_lanza(self):
        origen = self.ruta('malo.xlsx')
        with open(origen, 'wb') as f:
            f.write(b'no soy un zip')

        copia, temp_dir = parchear_tamano_papel(origen, 'Folio')

        self.assertEqual(copia, origen)
        self.assertIsNone(temp_dir)

    def test_papel_vacio_devuelve_el_original(self):
        origen = crear_xlsx(self.ruta('origen.xlsx'), ['<pageSetup paperSize="9"/>'])

        copia, temp_dir = parchear_tamano_papel(origen, '')

        self.assertEqual(copia, origen)
        self.assertIsNone(temp_dir)

    def test_formato_binario_no_parcheable(self):
        origen = self.ruta('viejo.xls')
        with open(origen, 'wb') as f:
            f.write(b'\xd0\xcf\x11\xe0')

        copia, temp_dir = parchear_tamano_papel(origen, 'Folio')

        self.assertEqual(copia, origen)
        self.assertIsNone(temp_dir)

    # --- márgenes superior/inferior -------------------------------------

    def test_aplica_margenes_sin_tocar_el_original(self):
        origen = crear_xlsx(self.ruta('margenes.xlsx'), [
            '<pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75"'
            ' header="0.3" footer="0.3"/><pageSetup paperSize="9"/>',
        ])

        copia, temp_dir = parchear_tamano_papel(origen, 'A4')

        self.assertEqual(margenes_xlsx(copia), (0.4, 0.4, 0.7))
        # El original conserva sus márgenes y su papel
        self.assertEqual(margenes_xlsx(origen), (0.75, 0.75, 0.7))
        self.assertEqual(paper_sizes(origen), ['9'])
        if temp_dir:
            shutil.rmtree(temp_dir, ignore_errors=True)

    def test_margenes_preservan_laterales_y_encabezado(self):
        origen = crear_xlsx(self.ruta('lateral.xlsx'), [
            '<pageMargins left="1.2" right="1.3" top="0.75" bottom="0.75"'
            ' header="0.3" footer="0.3"/><pageSetup paperSize="9"/>',
        ])

        copia, temp_dir = parchear_tamano_papel(origen, 'A4')

        with zipfile.ZipFile(copia) as z:
            etiqueta = re.search(
                r'<pageMargins[^>]*/>',
                z.read('xl/worksheets/sheet1.xml').decode('utf-8', 'replace'),
            ).group(0)
        self.assertIn('left="1.2"', etiqueta)
        self.assertIn('right="1.3"', etiqueta)
        self.assertIn('header="0.3"', etiqueta)
        self.assertIn('top="0.4"', etiqueta)
        self.assertIn('bottom="0.4"', etiqueta)
        if temp_dir:
            shutil.rmtree(temp_dir, ignore_errors=True)

    def test_inserta_page_margins_si_falta(self):
        origen = crear_xlsx(self.ruta('sin_margenes.xlsx'), ['<pageSetup paperSize="9"/>'])

        copia, temp_dir = parchear_tamano_papel(origen, 'A4')

        self.assertIsNotNone(margenes_xlsx(copia))
        self.assertEqual(margenes_xlsx(copia)[:2], (0.4, 0.4))
        if temp_dir:
            shutil.rmtree(temp_dir, ignore_errors=True)

    @override_settings(MARGEN_SUPERIOR_PDF_CONVERSION=None,
                       MARGEN_INFERIOR_PDF_CONVERSION=None)
    def test_margenes_desactivados_respetan_los_del_archivo(self):
        origen = crear_xlsx(self.ruta('m_origen.xlsx'), [
            '<pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75"'
            ' header="0.3" footer="0.3"/><pageSetup paperSize="9"/>',
        ])

        copia, temp_dir = parchear_tamano_papel(origen, 'A4')

        self.assertEqual(margenes_xlsx(copia), (0.75, 0.75, 0.7))
        self.assertEqual(paper_sizes(copia), ['9'])
        if temp_dir:
            shutil.rmtree(temp_dir, ignore_errors=True)


class ParcheoDocxTest(BaseTemporalTest):

    def test_retrato_gana_alto(self):
        origen = crear_docx(self.ruta('retrato.docx'), 12240, 15840)

        copia, temp_dir = parchear_tamano_papel(origen, 'Folio')

        self.assertEqual(pg_sz_docx(copia), (12240, 18720, None))
        self.assertEqual(pg_sz_docx(origen), (12240, 15840, None))
        if temp_dir:
            shutil.rmtree(temp_dir, ignore_errors=True)

    def test_apaisado_intercambia_medidas(self):
        origen = crear_docx(self.ruta('apaisado.docx'), 15840, 12240, orient='landscape')

        copia, temp_dir = parchear_tamano_papel(origen, 'Folio')

        self.assertEqual(pg_sz_docx(copia), (18720, 12240, 'landscape'))
        if temp_dir:
            shutil.rmtree(temp_dir, ignore_errors=True)

    def test_insere_pg_sz_si_falta(self):
        origen = self.ruta('sin_pgsz.docx')
        xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<w:document %s><w:body><w:p/>'
            '<w:sectPr><w:pgMar w:top="1440" w:right="1440" '
            'w:bottom="1440" w:left="1440"/></w:sectPr>'
            '</w:body></w:document>'
        ) % NS_W
        with zipfile.ZipFile(origen, 'w') as z:
            z.writestr('word/document.xml', xml)

        copia, temp_dir = parchear_tamano_papel(origen, 'Folio')

        self.assertEqual(pg_sz_docx(copia), (12240, 18720, None))
        if temp_dir:
            shutil.rmtree(temp_dir, ignore_errors=True)

    # --- márgenes superior/inferior -------------------------------------

    def test_aplica_margenes_superior_e_inferior(self):
        origen = crear_docx(self.ruta('margenes.docx'), 12240, 15840)

        copia, temp_dir = parchear_tamano_papel(origen, 'A4')

        # 0.4" = 576 twips; el lado izquierdo se respeta
        self.assertEqual(pg_mar_docx(copia), (576, 576, 1440))
        self.assertEqual(pg_mar_docx(origen), (1440, 1440, 1440))
        self.assertEqual(pg_sz_docx(copia)[2], None)
        if temp_dir:
            shutil.rmtree(temp_dir, ignore_errors=True)

    @override_settings(MARGEN_SUPERIOR_PDF_CONVERSION=None,
                       MARGEN_INFERIOR_PDF_CONVERSION=None)
    def test_margenes_desactivados_respetan_los_del_docx(self):
        origen = crear_docx(self.ruta('m_origen.docx'), 12240, 15840)

        copia, temp_dir = parchear_tamano_papel(origen, 'A4')

        self.assertEqual(pg_mar_docx(copia), (1440, 1440, 1440))
        if temp_dir:
            shutil.rmtree(temp_dir, ignore_errors=True)


# =============================================================================
# SANEO DE LA COPIA: salto de página y conservación de logos
# =============================================================================

class LogoSeccionTest(BaseTemporalTest):
    """Cada sección lleva su logo y el saneo no debe retirar ninguno."""

    def test_conserva_los_dos_logos(self):
        origen = xlsx_secciones_con_logo(self.ruta('logos.xlsx'))
        self.assertEqual(sorted(anclas_dibujo(origen)),
                         ['oneCellAnchor', 'twoCellAnchor'])

        copia, temp_dir = parchear_tamano_papel(origen, 'A4')

        # El logo de la última página (oneCell de la 2ª sección) sigue ahí
        self.assertEqual(sorted(anclas_dibujo(copia)),
                         ['oneCellAnchor', 'twoCellAnchor'])
        self.assertEqual(filas_con_salto(copia), [29])
        if temp_dir:
            shutil.rmtree(temp_dir, ignore_errors=True)

    def test_el_original_no_se_modifica(self):
        origen = xlsx_secciones_con_logo(self.ruta('logos_intacto.xlsx'))
        antes = anclas_dibujo(origen)

        copia, temp_dir = parchear_tamano_papel(origen, 'A4')

        self.assertEqual(anclas_dibujo(origen), antes)
        self.assertEqual(sorted(anclas_dibujo(copia)),
                         ['oneCellAnchor', 'twoCellAnchor'])
        if temp_dir:
            shutil.rmtree(temp_dir, ignore_errors=True)

    def test_imagen_unica_no_se_toca(self):
        """Sin logo de cabecera, el de la 2ª sección se queda igual."""
        origen = xlsx_secciones_con_logo(
            self.ruta('unica.xlsx'), encabezado=False,
        )
        self.assertEqual(anclas_dibujo(origen), ['oneCellAnchor'])

        copia, temp_dir = parchear_tamano_papel(origen, 'A4')

        self.assertEqual(anclas_dibujo(copia), ['oneCellAnchor'])
        if temp_dir:
            shutil.rmtree(temp_dir, ignore_errors=True)

    def test_dibujo_corrupto_no_lanza(self):
        origen = self.ruta('dibujo_malo.xlsx')
        with zipfile.ZipFile(origen, 'w', zipfile.ZIP_DEFLATED) as z:
            z.writestr('[Content_Types].xml', '<?xml version="1.0"?><Types/>')
            z.writestr(
                'xl/worksheets/sheet1.xml',
                '<?xml version="1.0"?><worksheet %s><sheetData/></worksheet>' % NS_XLSX,
            )
            z.writestr('xl/drawings/drawing1.xml', 'no soy xml')

        copia, temp_dir = parchear_tamano_papel(origen, 'A4')

        self.assertTrue(os.path.exists(copia))
        with zipfile.ZipFile(copia) as z:
            self.assertEqual(z.read('xl/drawings/drawing1.xml'), b'no soy xml')
        if temp_dir:
            shutil.rmtree(temp_dir, ignore_errors=True)


class SaltoTituloTest(BaseTemporalTest):
    """El título repetido (2ª sección) debe abrir página en la copia."""

    def test_salta_antes_del_titulo_repetido(self):
        origen = xlsx_con_titulo_repetido(self.ruta('titulo.xlsx'), repite_en=30)

        copia, temp_dir = parchear_tamano_papel(origen, 'A4')

        # LibreOffice pone en `id` la última fila de la página anterior
        self.assertEqual(filas_con_salto(copia), [29])
        self.assertEqual(filas_con_salto(origen), [])
        if temp_dir:
            shutil.rmtree(temp_dir, ignore_errors=True)

    def test_sin_titulo_repetido_no_hay_salto(self):
        origen = xlsx_con_titulo_repetido(self.ruta('sin.xlsx'), repite_en=None)

        copia, temp_dir = parchear_tamano_papel(origen, 'A4')

        self.assertEqual(filas_con_salto(copia), [])
        if temp_dir:
            shutil.rmtree(temp_dir, ignore_errors=True)

    def test_solo_salta_en_la_primera_repeticion(self):
        origen = xlsx_con_titulo_repetido(
            self.ruta('dos.xlsx'), repite_en=(30, 35),
        )

        copia, temp_dir = parchear_tamano_papel(origen, 'A4')

        self.assertEqual(filas_con_salto(copia), [29])
        if temp_dir:
            shutil.rmtree(temp_dir, ignore_errors=True)

    def test_no_duplica_un_salto_existente(self):
        origen = xlsx_con_titulo_repetido(
            self.ruta('previo.xlsx'), repite_en=30, salto_previo=29,
        )

        copia, temp_dir = parchear_tamano_papel(origen, 'A4')

        self.assertEqual(filas_con_salto(copia), [29])
        if temp_dir:
            shutil.rmtree(temp_dir, ignore_errors=True)

    def test_el_hoja_del_formulario_no_salta(self):
        """La hoja sin título repetido (formulario) queda sin <rowBreaks>."""
        origen = xlsx_con_titulo_repetido(self.ruta('dos_hojas.xlsx'), repite_en=30)
        with zipfile.ZipFile(origen) as z:
            partes = {n: z.read(n) for n in z.namelist()}
        partes['xl/worksheets/sheet2.xml'] = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<worksheet %s><sheetData><row r="1"><c r="A1" t="inlineStr">'
            '<is><t>Formulario</t></is></c></row></sheetData></worksheet>'
        ) % NS_XLSX
        with zipfile.ZipFile(origen, 'w', zipfile.ZIP_DEFLATED) as z:
            for nombre, contenido in partes.items():
                z.writestr(nombre, contenido)

        copia, temp_dir = parchear_tamano_papel(origen, 'A4')

        self.assertEqual(filas_con_salto(copia), [29])
        self.assertEqual(filas_con_salto(copia, 'xl/worksheets/sheet2.xml'), [])
        if temp_dir:
            shutil.rmtree(temp_dir, ignore_errors=True)


# =============================================================================
# CONVERSIÓN (LibreOffice mockeado)
# =============================================================================

class ConvertirOfficePdfTest(BaseTemporalTest):

    @staticmethod
    def _run_falso(cmd, **kwargs):
        """LibreOffice simulado: escribe el PDF en el --outdir indicado."""
        entrada = cmd[-1]
        outdir = cmd[cmd.index('--outdir') + 1]
        nombre = os.path.splitext(os.path.basename(entrada))[0] + '.pdf'
        with open(os.path.join(outdir, nombre), 'wb') as f:
            f.write(b'%PDF-1.4 fake')
        return subprocess.CompletedProcess(cmd, 0, stdout='', stderr='')

    def test_convierte_sobre_la_copia_parcheada(self):
        origen = crear_xlsx(self.ruta('origen.xlsx'), ['<pageSetup paperSize="9"/>'])
        captura = {}

        def run(cmd, **kwargs):
            captura['entrada'] = cmd[-1]
            captura['tallas'] = paper_sizes(cmd[-1])
            return self._run_falso(cmd, **kwargs)

        with patch('documentos.services.conversion_pdf.subprocess.run', side_effect=run):
            pdf = convertir_office_a_pdf(origen, papel='Folio', outdir=self.dir)

        self.assertEqual(pdf, os.path.join(self.dir, 'origen.pdf'))
        self.assertTrue(os.path.exists(pdf))
        self.assertNotEqual(captura['entrada'], origen)
        self.assertEqual(captura['tallas'], ['14'])
        # El archivo subido queda intacto
        self.assertEqual(paper_sizes(origen), ['9'])

    def test_sin_papel_convierte_el_original(self):
        origen = crear_xlsx(self.ruta('origen.xlsx'), ['<pageSetup paperSize="9"/>'])
        captura = {}

        def run(cmd, **kwargs):
            captura['entrada'] = cmd[-1]
            return self._run_falso(cmd, **kwargs)

        with patch('documentos.services.conversion_pdf.subprocess.run', side_effect=run):
            convertir_office_a_pdf(origen, papel='', outdir=self.dir)

        self.assertEqual(captura['entrada'], origen)

    def test_limpia_la_copia_temporal(self):
        origen = crear_xlsx(self.ruta('origen.xlsx'), ['<pageSetup paperSize="9"/>'])
        captura = {}

        def run(cmd, **kwargs):
            captura['entrada'] = cmd[-1]
            return self._run_falso(cmd, **kwargs)

        with patch('documentos.services.conversion_pdf.subprocess.run', side_effect=run):
            convertir_office_a_pdf(origen, papel='Folio', outdir=self.dir)

        self.assertFalse(os.path.exists(os.path.dirname(captura['entrada'])))

    def test_formato_no_convertible(self):
        origen = self.ruta('notas.txt')
        with open(origen, 'w') as f:
            f.write('hola')

        with patch('documentos.services.conversion_pdf.subprocess.run') as run:
            self.assertIsNone(convertir_office_a_pdf(origen, outdir=self.dir))
        run.assert_not_called()

    def test_error_de_libreoffice(self):
        origen = crear_xlsx(self.ruta('origen.xlsx'), ['<pageSetup paperSize="9"/>'])

        with patch(
            'documentos.services.conversion_pdf.subprocess.run',
            return_value=subprocess.CompletedProcess([], 1, stdout='', stderr='boom'),
        ):
            self.assertIsNone(convertir_office_a_pdf(origen, outdir=self.dir))


# =============================================================================
# CONVERSIÓN REAL (LibreOffice)
# =============================================================================

@unittest.skipUnless(LIBREOFFICE, 'LibreOffice no instalado')
class ConversionRealLibreOfficeTest(BaseTemporalTest):
    """Comprueba de verdad que el papel y los márgenes llegan al PDF."""

    def test_folio_da_8_5_por_13_pulgadas(self):
        origen = xlsx_real(self.ruta('origen.xlsx'), paper_size=9)

        pdf = convertir_office_a_pdf(origen, papel='Folio', outdir=self.dir)

        self.assertTrue(pdf, 'LibreOffice no generó el PDF')
        ancho, alto = primera_mediabox(pdf)
        self.assertEqual((round(ancho), round(alto)), (612, 936))
        # El original sigue en A4
        self.assertEqual(paper_sizes(origen), ['9'])

    def test_papel_vacio_respeta_el_del_archivo(self):
        origen = xlsx_real(self.ruta('origen.xlsx'), paper_size=9)

        pdf = convertir_office_a_pdf(origen, papel='', outdir=self.dir)

        self.assertTrue(pdf, 'LibreOffice no generó el PDF')
        ancho, alto = primera_mediabox(pdf)
        self.assertEqual((round(ancho), round(alto)), (595, 842))

    def test_a4_es_el_papel_por_defecto(self):
        origen = xlsx_real(self.ruta('defecto.xlsx'), paper_size=9)

        pdf = convertir_office_a_pdf(origen, outdir=self.dir)

        self.assertTrue(pdf, 'LibreOffice no generó el PDF')
        ancho, alto = primera_mediabox(pdf)
        self.assertEqual((round(ancho), round(alto)), (595, 842))
        # El original sigue declarando A4 y sus márgenes
        self.assertEqual(paper_sizes(origen), ['9'])

    def test_margenes_reducidos_meten_el_contenido_en_una_pagina(self):
        """50 filas de 15pt = 750pt.

        Con los márgenes originales (0.75") el área útil es 734pt y la última
        fila se va a la página 2; con los nuestros (0.4") son 784pt y todo
        cabe en una sola página.
        """
        origen = xlsx_de_filas(self.ruta('lleno.xlsx'), filas=50)

        con_margenes = convertir_office_a_pdf(origen, outdir=self.ruta('con'))
        sin_margenes = convertir_office_a_pdf(
            origen, papel='A4', outdir=self.ruta('sin'),
            margenes={'top': None, 'bottom': None},
        )

        self.assertTrue(con_margenes and sin_margenes, 'LibreOffice falló')
        self.assertEqual(contar_paginas(con_margenes), 1)
        self.assertEqual(contar_paginas(sin_margenes), 2)
        # El original sigue con sus márgenes
        self.assertEqual(margenes_xlsx(origen), (0.75, 0.75, 0.75))


@unittest.skipUnless(LIBREOFFICE, 'LibreOffice no instalado')
class SaneoRealLibreOfficeTest(BaseTemporalTest):
    """Comprueba el saneo en el PDF que LibreOffice genera de verdad."""

    def test_el_titulo_repetido_abre_pagina(self):
        origen = xlsx_con_titulo_repetido(self.ruta('titulo.xlsx'))

        with patch(
            'documentos.services.conversion_pdf._fila_titulo_repetido',
            return_value=None,
        ):
            sin_salto = convertir_office_a_pdf(origen, outdir=self.ruta('sin'))
        con_salto = convertir_office_a_pdf(origen, outdir=self.ruta('con'))

        self.assertTrue(sin_salto and con_salto, 'LibreOffice no generó el PDF')
        self.assertEqual(contar_paginas(sin_salto), 1)
        self.assertEqual(contar_paginas(con_salto), 2)
        segunda = paginas_texto(con_salto)[1].strip()
        self.assertTrue(
            segunda.startswith('TITULO DE PRUEBA'),
            'La segunda página no empieza por el título: %r' % segunda[:120],
        )

    def test_el_logo_de_la_ultima_pagina_se_dibuja(self):
        origen = xlsx_secciones_con_logo(self.ruta('logos.xlsx'))

        pdf = convertir_office_a_pdf(origen, outdir=self.ruta('con'))

        self.assertTrue(pdf, 'LibreOffice no generó el PDF')
        self.assertEqual(contar_paginas(pdf), 2)

        dibujos = dibujos_pdf(pdf)
        # la 2ª sección abre la última página y su logo va con ella: borrarlo
        # "por duplicado" dejaba esa página sin logo
        ultima = [d for d in dibujos if d[0] == 2]
        self.assertEqual(
            len(ultima), 1, 'la última página debe llevar su logo: %r' % (dibujos,),
        )
        for _, matriz, _ in dibujos:
            self.assertGreaterEqual(
                float(matriz.split()[-1]), 0,
                'logo recortado bajo el borde inferior: %s' % matriz,
            )

        segunda = paginas_texto(pdf)[1].strip()
        self.assertTrue(
            segunda.startswith('TITULO DE PRUEBA'),
            'La segunda página no empieza por el título: %r' % segunda[:120],
        )


# =============================================================================
# TAREA CELERY
# =============================================================================

class TareaConversionTest(BaseDocumentosTest):

    @classmethod
    def setUpClass(cls):
        cls._media_dir = tempfile.mkdtemp(prefix='polidoc-conv-')
        cls._media_override = override_settings(MEDIA_ROOT=cls._media_dir)
        cls._media_override.enable()
        try:
            super().setUpClass()
        except Exception:
            cls._media_override.disable()
            shutil.rmtree(cls._media_dir, ignore_errors=True)
            raise
        cls.addClassCleanup(cls._media_override.disable)
        cls.addClassCleanup(shutil.rmtree, cls._media_dir, ignore_errors=True)

    def setUp(self):
        super().setUp()
        self.dir = tempfile.mkdtemp(prefix='polidoc-conv-test-')
        self.addCleanup(shutil.rmtree, self.dir, ignore_errors=True)

    def ruta(self, nombre):
        return os.path.join(self.dir, nombre)

    @classmethod
    def _crear_documento(cls, codigo, contenido=b'xlsx', nombre='origen.xlsx'):
        return Documento.objects.create(
            tipo=cls.tipo_formulario, nombre='Documento de conversión',
            codigo=codigo, estado='activo', propietario=cls.user_trabajador,
            departamento=cls.depto_ti, tipo_archivo='xlsx',
            archivo=SimpleUploadedFile(nombre, contenido),
        )

    @override_settings(PAPEL_PDF_CONVERSION='Folio')
    def test_tarea_pasa_el_papel_configurado(self):
        doc = self._crear_documento('TI-FOR-CONV-1')
        llamada = {}

        def falso_convertir(archivo, papel=None, outdir=None, timeout=180):
            llamada['papel'] = papel
            llamada['outdir'] = outdir
            pdf = os.path.join(outdir, os.path.splitext(os.path.basename(archivo))[0] + '.pdf')
            with open(pdf, 'wb') as f:
                f.write(b'%PDF-1.4 fake')
            return pdf

        with patch(
            'documentos.services.conversion_pdf.convertir_office_a_pdf',
            side_effect=falso_convertir,
        ) as mock:
            convertir_documento_a_pdf(str(doc.uuid))

        mock.assert_called_once()
        self.assertEqual(llamada['papel'], 'Folio')
        self.assertEqual(llamada['outdir'], os.path.dirname(doc.archivo.path))
        doc.refresh_from_db()
        self.assertTrue(doc.archivo_pdf)
        esperado = os.path.splitext(doc.archivo.name)[0] + '.pdf'
        self.assertEqual(doc.archivo_pdf.name, esperado)
        self.assertFalse(os.path.isabs(doc.archivo_pdf.name))

    @unittest.skipUnless(LIBREOFFICE, 'LibreOffice no instalado')
    @override_settings(PAPEL_PDF_CONVERSION='Folio')
    def test_tarea_real_genera_un_pdf_folio(self):
        contenido = xlsx_real(self.ruta('origen-real.xlsx'), paper_size=9)
        with open(contenido, 'rb') as f:
            datos = f.read()
        doc = self._crear_documento('TI-FOR-CONV-2', contenido=datos, nombre='real.xlsx')

        convertir_documento_a_pdf(str(doc.uuid))

        doc.refresh_from_db()
        self.assertTrue(doc.archivo_pdf, 'No se guardó archivo_pdf')
        pdf = os.path.join(str(self._media_dir), doc.archivo_pdf.name)
        self.assertTrue(os.path.exists(pdf), doc.archivo_pdf.name)
        ancho, alto = primera_mediabox(pdf)
        self.assertEqual((round(ancho), round(alto)), (612, 936))
        # El XLSX subido no se tocó: sigue declarando A4
        self.assertEqual(paper_sizes(doc.archivo.path), ['9'])

    # --- Un solo conversor para documentos, planillas y exportación -------

    def test_convertir_a_pdf_sync_delega_en_el_conversor_comun(self):
        origen = os.path.join(self._media_dir, 'planillas', 'origen.xlsx')
        os.makedirs(os.path.dirname(origen), exist_ok=True)
        xlsx_real(origen)
        llamada = {}

        def falso_convertir(archivo, papel=None, outdir=None, timeout=180,
                            margenes=None):
            llamada.update(archivo=archivo, papel=papel, outdir=outdir,
                           margenes=margenes)
            pdf = os.path.join(outdir, 'origen.pdf')
            with open(pdf, 'wb') as f:
                f.write(b'%PDF-1.4 fake')
            return pdf

        with patch(
            'documentos.services.conversion_pdf.convertir_office_a_pdf',
            side_effect=falso_convertir,
        ) as mock:
            ruta = convertir_a_pdf_sync(origen)

        mock.assert_called_once()
        self.assertEqual(llamada['archivo'], origen)
        # el papel y los márgenes los decide settings dentro del conversor
        self.assertIsNone(llamada['papel'])
        self.assertEqual(llamada['outdir'], os.path.dirname(origen))
        # la ruta devuelta sigue siendo relativa a MEDIA_ROOT
        self.assertEqual(ruta, os.path.join('planillas', 'origen.pdf'))

    def test_convertir_a_pdf_sync_sin_libreoffice_devuelve_none(self):
        origen = os.path.join(self._media_dir, 'planillas', 'otro.xlsx')
        os.makedirs(os.path.dirname(origen), exist_ok=True)
        xlsx_real(origen)

        with patch(
            'documentos.services.conversion_pdf.convertir_office_a_pdf',
            return_value=None,
        ):
            self.assertIsNone(convertir_a_pdf_sync(origen))

    def test_convertir_planilla_a_pdf_usa_el_mismo_conversor(self):
        doc = self._crear_documento('TI-FOR-CONV-3')
        planilla = PlanillaCambios.objects.create(
            codigo='PC-CONV-1', documento=doc, tipo_cambio='modificacion',
            justificacion='Vista previa', responsable=self.user_jefe_ti,
            version_destino='2.0', estado='en_revision',
            archivo_propuesto=SimpleUploadedFile('propuesto.xlsx', b'xlsx'),
        )

        with patch(
            'documentos.services.conversion_pdf.convertir_office_a_pdf',
            side_effect=_pdf_falso,
        ) as mock:
            resultado = convertir_planilla_a_pdf(planilla.id)

        mock.assert_called_once()
        self.assertTrue(resultado['success'], resultado)
        planilla.refresh_from_db()
        self.assertTrue(planilla.archivo_propuesto_pdf)
        self.assertFalse(os.path.isabs(planilla.archivo_propuesto_pdf.name))


# =============================================================================
# REGENERACIÓN AUTOMÁTICA (ruta fantasma en la BD)
# =============================================================================

def _pdf_falso(archivo, papel=None, outdir=None, timeout=180):
    """LibreOffice simulado para services.documentos_pdf."""
    nombre = os.path.splitext(os.path.basename(archivo))[0] + '.pdf'
    ruta = os.path.join(outdir, nombre)
    with open(ruta, 'wb') as f:
        f.write(b'%PDF-1.4 fake')
    return ruta


class BaseRegeneracionTest(BaseDocumentosTest):
    """MEDIA_ROOT temporal + helpers para documentos con PDF fantasma."""

    @classmethod
    def setUpClass(cls):
        cls._media_dir = tempfile.mkdtemp(prefix='polidoc-regen-')
        cls._media_override = override_settings(MEDIA_ROOT=cls._media_dir)
        cls._media_override.enable()
        try:
            super().setUpClass()
        except Exception:
            cls._media_override.disable()
            shutil.rmtree(cls._media_dir, ignore_errors=True)
            raise
        cls.addClassCleanup(cls._media_override.disable)
        cls.addClassCleanup(shutil.rmtree, cls._media_dir, ignore_errors=True)

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        with tempfile.TemporaryDirectory() as tmp:
            origen = xlsx_real(os.path.join(tmp, 'origen.xlsx'), paper_size=9)
            with open(origen, 'rb') as fh:
                cls._xlsx_bytes = fh.read()

    def _doc(self, codigo, fantasma=False, tipo_archivo='xlsx'):
        if tipo_archivo == 'pdf':
            contenido = b'%PDF-1.4 contenido original'
            nombre = 'origen.pdf'
        else:
            contenido = self._xlsx_bytes
            nombre = 'origen.xlsx'
        doc = Documento.objects.create(
            tipo=self.tipo_formulario, nombre='Documento de regeneración',
            codigo=codigo, estado='activo', propietario=self.user_trabajador,
            departamento=self.depto_ti, clasificacion='interno',
            version_actual='1.0', tipo_archivo=tipo_archivo,
            archivo=SimpleUploadedFile(nombre, contenido),
        )
        self.addCleanup(self._borrar_ficheros, doc)
        if fantasma:
            doc.archivo_pdf = 'documentos/fantasma/%s.pdf' % codigo
            doc.save(update_fields=['archivo_pdf'])
        return doc

    @staticmethod
    def _borrar_ficheros(doc):
        for campo in (doc.archivo, doc.archivo_pdf):
            try:
                if campo and campo.name:
                    campo.storage.delete(campo.name)
            except Exception:
                pass

    def _url(self, doc, vista):
        return f'/documentos/{doc.tipo_id}/{doc.uuid}/{vista}'

    def _doc_con_pdf_real(self, codigo):
        """Documento con un PDF realmente presente en disco."""
        doc = self._doc(codigo)
        pdf_path = os.path.join(os.path.dirname(doc.archivo.path), 'pre.pdf')
        os.makedirs(os.path.dirname(pdf_path), exist_ok=True)
        with open(pdf_path, 'wb') as f:
            f.write(b'%PDF-1.4 pre')
        rel = pdf_path[len(str(self._media_dir)):].lstrip(os.sep)
        doc.archivo_pdf = rel
        doc.save(update_fields=['archivo_pdf'])
        return doc


class PdfDisponibleTest(BaseRegeneracionTest):

    def test_sin_campo_no_esta_disponible(self):
        doc = self._doc('TI-FOR-PDF-1')
        self.assertFalse(pdf_disponible(doc))

    def test_ruta_fantasma_no_esta_disponible(self):
        doc = self._doc('TI-FOR-PDF-2', fantasma=True)
        self.assertFalse(pdf_disponible(doc))

    def test_fichero_presente_si_esta_disponible(self):
        doc = self._doc_con_pdf_real('TI-FOR-PDF-3')
        self.assertTrue(pdf_disponible(doc))

    def test_limpia_puntero_fantasma(self):
        doc = self._doc('TI-FOR-PDF-4', fantasma=True)

        self.assertTrue(limpiar_pdf_inexistente(doc))

        doc.refresh_from_db()
        self.assertFalse(doc.archivo_pdf)
        # Ya no hay nada que limpiar
        self.assertFalse(limpiar_pdf_inexistente(doc))

    def test_limpia_no_toca_un_pdf_valido(self):
        doc = self._doc_con_pdf_real('TI-FOR-PDF-5')

        self.assertFalse(limpiar_pdf_inexistente(doc))

        doc.refresh_from_db()
        self.assertTrue(doc.archivo_pdf)

    def test_asegurar_pdf_regenera_el_pdf(self):
        doc = self._doc('TI-FOR-PDF-6', fantasma=True)

        with patch('documentos.services.conversion_pdf.convertir_office_a_pdf',
                   side_effect=_pdf_falso) as mock:
            self.assertTrue(asegurar_pdf(doc))

        mock.assert_called_once()
        doc.refresh_from_db()
        self.assertTrue(pdf_disponible(doc))

    def test_asegurar_pdf_no_reconvierte_si_ya_existe(self):
        doc = self._doc_con_pdf_real('TI-FOR-PDF-7')

        with patch('documentos.services.conversion_pdf.convertir_office_a_pdf',
                   side_effect=_pdf_falso) as mock:
            self.assertTrue(asegurar_pdf(doc))

        mock.assert_not_called()

    def test_asegurar_pdf_no_intent_si_no_es_convertible(self):
        doc = self._doc('TI-FOR-PDF-8', tipo_archivo='pdf')

        with patch('documentos.services.conversion_pdf.convertir_office_a_pdf',
                   side_effect=_pdf_falso) as mock:
            self.assertFalse(asegurar_pdf(doc))

        mock.assert_not_called()

    def test_asegurar_pdf_sin_origen_no_lanza(self):
        doc = self._doc('TI-FOR-PDF-9', fantasma=True)
        doc.archivo.delete(save=False)

        with patch('documentos.services.conversion_pdf.convertir_office_a_pdf') as mock:
            self.assertFalse(asegurar_pdf(doc))

        mock.assert_not_called()
        doc.refresh_from_db()
        self.assertFalse(doc.archivo_pdf)

    def test_convertir_documento_sync_falla_limpio(self):
        doc = self._doc('TI-FOR-PDF-10', fantasma=True)

        with patch('documentos.services.conversion_pdf.convertir_office_a_pdf',
                   return_value=None):
            self.assertFalse(convertir_documento_sync(doc))

        doc.refresh_from_db()
        self.assertFalse(doc.archivo_pdf)

    def test_archivo_existe_con_campo_vacio(self):
        doc = self._doc('TI-FOR-PDF-11')
        self.assertFalse(archivo_existe(None))
        self.assertTrue(archivo_existe(doc.archivo))


class RegeneracionVistasTest(BaseRegeneracionTest):
    """Las vistas no deben dar 500 si el PDF desapareció del disco."""

    def setUp(self):
        super().setUp()
        self.client = Client()

    def test_ver_archivo_regenera_el_pdf_faltante(self):
        doc = self._doc('TI-FOR-V-1', fantasma=True)
        self.client.login(username='trab01', password='testpass123')

        with patch('documentos.services.conversion_pdf.convertir_office_a_pdf',
                   side_effect=_pdf_falso):
            resp = self.client.get(self._url(doc, 'ver-archivo/'))

        self.assertEqual(resp.status_code, 200)
        self.assertEqual(b''.join(resp.streaming_content), b'%PDF-1.4 fake')
        doc.refresh_from_db()
        self.assertTrue(pdf_disponible(doc))

    def test_ver_archivo_sin_conversion_no_da_500(self):
        doc = self._doc('TI-FOR-V-2', fantasma=True)
        self.client.login(username='trab01', password='testpass123')

        with patch('documentos.services.conversion_pdf.convertir_office_a_pdf',
                   return_value=None):
            resp = self.client.get(self._url(doc, 'ver-archivo/'))

        self.assertEqual(resp.status_code, 302)
        doc.refresh_from_db()
        self.assertFalse(doc.archivo_pdf)

    def test_ver_archivo_pdf_presente_se_sirve(self):
        doc = self._doc_con_pdf_real('TI-FOR-V-3')
        self.client.login(username='trab01', password='testpass123')

        with patch('documentos.services.conversion_pdf.convertir_office_a_pdf') as mock:
            resp = self.client.get(self._url(doc, 'ver-archivo/'))

        self.assertEqual(resp.status_code, 200)
        self.assertEqual(b''.join(resp.streaming_content), b'%PDF-1.4 pre')
        mock.assert_not_called()

    def test_ver_archivo_origen_ausente_no_da_500(self):
        doc = self._doc('TI-FOR-V-4')
        doc.archivo.delete(save=False)
        self.client.login(username='trab01', password='testpass123')

        with patch('documentos.services.conversion_pdf.convertir_office_a_pdf') as mock:
            resp = self.client.get(self._url(doc, 'ver-archivo/'))

        self.assertEqual(resp.status_code, 302)
        mock.assert_not_called()

    def test_ver_archivo_no_toca_el_archivo_original(self):
        doc = self._doc('TI-FOR-V-5', fantasma=True)
        self.client.login(username='trab01', password='testpass123')

        with patch('documentos.services.conversion_pdf.convertir_office_a_pdf',
                   side_effect=_pdf_falso):
            self.client.get(self._url(doc, 'ver-archivo/'))

        self.assertTrue(os.path.exists(doc.archivo.path))
        self.assertEqual(paper_sizes(doc.archivo.path), ['9'])

    def test_detalle_regenera_el_pdf_faltante(self):
        doc = self._doc('TI-FOR-V-6', fantasma=True)
        self.client.login(username='trab01', password='testpass123')

        with patch('documentos.services.conversion_pdf.convertir_office_a_pdf',
                   side_effect=_pdf_falso):
            resp = self.client.get(self._url(doc, ''))

        self.assertEqual(resp.status_code, 200)
        doc.refresh_from_db()
        self.assertTrue(pdf_disponible(doc))

    def test_detalle_sin_conversion_no_da_500(self):
        doc = self._doc('TI-FOR-V-7', fantasma=True)
        self.client.login(username='trab01', password='testpass123')

        with patch('documentos.services.conversion_pdf.convertir_office_a_pdf',
                   return_value=None):
            resp = self.client.get(self._url(doc, ''))

        self.assertEqual(resp.status_code, 200)
        doc.refresh_from_db()
        self.assertFalse(doc.archivo_pdf)

    def test_estado_conversion_espelja_el_disco(self):
        doc_fantasma = self._doc('TI-FOR-V-8', fantasma=True)
        doc_ok = self._doc_con_pdf_real('TI-FOR-V-9')
        self.client.login(username='trab01', password='testpass123')

        with patch('documentos.services.conversion_pdf.convertir_office_a_pdf') as mock:
            fantasma = self.client.get(self._url(doc_fantasma, 'estado-conversion/'))
            listo = self.client.get(self._url(doc_ok, 'estado-conversion/'))

        self.assertEqual(fantasma.json()['listo'], False)
        self.assertEqual(listo.json()['listo'], True)
        mock.assert_not_called()

    def test_descargar_archivo_origen_ausente_no_da_500(self):
        doc = self._doc('TI-FOR-V-10')
        doc.archivo.delete(save=False)
        self.client.login(username='jefe_ti', password='testpass123')

        resp = self.client.get(self._url(doc, 'descargar/'))

        self.assertEqual(resp.status_code, 302)

    def test_descargar_archivo_origen_presente_se_sirve(self):
        doc = self._doc('TI-FOR-V-11')
        self.client.login(username='jefe_ti', password='testpass123')

        resp = self.client.get(self._url(doc, 'descargar/'))

        self.assertEqual(resp.status_code, 200)

    def test_aprobar_planilla_no_ofrece_un_pdf_inexistente(self):
        from django.urls import reverse

        doc = self._doc('TI-FOR-V-12', fantasma=True)
        planilla = PlanillaCambios.objects.create(
            codigo='PC-REGEN-1', documento=doc, tipo_cambio='modificacion',
            justificacion='Prueba de preview', responsable=self.user_jefe_ti,
            version_destino='2.0', estado='en_revision',
        )
        self.client.login(username='jefe_gg', password='testpass123')

        resp = self.client.get(reverse('documentos:aprobar_planilla',
                                       args=[planilla.uuid]))

        self.assertEqual(resp.status_code, 200)
        self.assertIsNone(resp.context['archivo_preview'])

    # --- Extremo a extremo con LibreOffice real (sin mocks) ----------------

    @unittest.skipUnless(LIBREOFFICE, 'LibreOffice no instalado')
    def test_ver_archivo_regenera_con_libreoffice_real(self):
        doc = self._doc('TI-FOR-V-RT', fantasma=True)
        self.client.login(username='trab01', password='testpass123')

        resp = self.client.get(self._url(doc, 'ver-archivo/'))

        self.assertEqual(resp.status_code, 200)
        cuerpo = b''.join(resp.streaming_content)
        self.assertTrue(cuerpo.startswith(b'%PDF'), 'No es un PDF')
        doc.refresh_from_db()
        self.assertTrue(pdf_disponible(doc))
        ancho, alto = primera_mediabox(doc.archivo_pdf.path)
        self.assertAlmostEqual(ancho, 595, delta=1)
        self.assertAlmostEqual(alto, 842, delta=1)
        # El archivo subido sigue intacto (A4 declarado, sin parchear)
        self.assertEqual(paper_sizes(doc.archivo.path), ['9'])

    @unittest.skipUnless(LIBREOFFICE, 'LibreOffice no instalado')
    def test_detalle_regenera_con_libreoffice_real(self):
        doc = self._doc('TI-FOR-V-RT2', fantasma=True)
        self.client.login(username='trab01', password='testpass123')

        resp = self.client.get(self._url(doc, ''))

        self.assertEqual(resp.status_code, 200)
        doc.refresh_from_db()
        self.assertTrue(pdf_disponible(doc))
        self.assertContains(resp, 'id="pdf-viewer"')
        self.assertNotContains(resp, 'id="conv-progress"')
        ancho, alto = primera_mediabox(doc.archivo_pdf.path)
        self.assertAlmostEqual(ancho, 595, delta=1)
        self.assertAlmostEqual(alto, 842, delta=1)
